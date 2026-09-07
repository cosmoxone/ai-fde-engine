"""
轻量 Markdown → DOCX 渲染器（python-docx）

覆盖交付物用到的语法子集：# / ## / ### 标题、段落、
- 列表、| 表格 |、``` 代码块、> 引用、**加粗** 行内标记。
非目标：完整 Markdown 规范。
"""

from __future__ import annotations

import io
import re
from typing import Optional


def markdown_to_docx(title: str, markdown: str) -> Optional[bytes]:
    """渲染为 DOCX 字节流；python-docx 未安装时抛 ImportError（由调用方降级）"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError as e:
        raise ImportError("python-docx 未安装: pip install python-docx") from e

    doc = Document()

    # 中文字体默认
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    # 封面标题
    h = doc.add_heading(title, level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    lines = markdown.split("\n")
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    def flush_table():
        if not table_buf:
            return
        rows = table_buf[:]
        table_buf.clear()
        header = rows[0]
        body = [r for r in rows[1:] if not all(set(c.strip()) <= {"-", " ", ":"} for c in r)] if len(rows) > 1 else []
        n_cols = max(len(r) for r in [header] + body)
        table = doc.add_table(rows=1, cols=n_cols)
        table.style = "Light Grid Accent 1"
        for j, text in enumerate(header[:n_cols]):
            cell = table.rows[0].cells[j]
            cell.text = _strip_inline(text)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
        for row in body:
            cells = table.add_row().cells
            for j, text in enumerate(row[:n_cols]):
                cells[j].text = _strip_inline(text)
                for p in cells[j].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)

    def flush_code():
        nonlocal code_buf
        if code_buf:
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_buf))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            code_buf = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # 表格
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            continue
        flush_table()

        # 代码块
        if stripped.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 3)
            doc.add_heading(_strip_inline(m.group(2)), level=level)
            i += 1
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_runs_with_bold(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        # 列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _add_runs_with_bold(p, stripped)
        i += 1

    flush_table()
    flush_code()

    # 页脚：导出信息
    from datetime import datetime

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"AI-FDE Engine 导出 · {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _strip_inline(text: str) -> str:
    """去掉行内 markdown 标记（表格单元格用纯文本）"""
    return _BOLD_RE.sub(r"\1", text).replace("`", "")


def _add_runs_with_bold(paragraph, text: str) -> None:
    """按 **加粗** 分段添加 runs"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()].replace("`", ""))
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:].replace("`", ""))
